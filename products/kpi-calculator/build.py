import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = sys.argv[1]
BLACK, GOLD, CREAM, INK, GREY, LINE = '0D0D0D', 'C9A84C', 'F3EFE4', '1A1A1A', '666666', 'D3CCB8'
F = 'Arial'
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT_FONT = Font(name=F, size=11, color='0000FF')
GBP = '£#,##0;(£#,##0);-'
NUM = '#,##0;(#,##0);-'
NUM1 = '#,##0.0;(#,##0.0);-'

wb = Workbook()


def title(ws, text, sub, last='F'):
    ws.merge_cells(f'B1:{last}1')
    ws['B1'] = text
    ws['B1'].font = Font(name=F, size=16, bold=True, color=GOLD)
    for col in 'ABCDEFGH'[:ord(last) - 64 + 1]:
        ws[f'{col}1'].fill = fill(BLACK)
        ws[f'{col}2'].fill = fill(BLACK)
    ws.row_dimensions[1].height = 34
    ws['B1'].alignment = Alignment(vertical='center')
    ws.merge_cells(f'B2:{last}2')
    ws['B2'] = sub
    ws['B2'].font = Font(name=F, size=9, italic=True, color='BBBBBB')
    ws.row_dimensions[2].height = 20
    ws['B2'].alignment = Alignment(vertical='top')
    for col in 'ABCDEFGH'[:ord(last) - 64 + 1]:
        ws[f'{col}3'].fill = fill(GOLD)
    ws.row_dimensions[3].height = 3


def section(ws, r, text, heads=None):
    ws[f'B{r}'] = text
    for col in 'BCDEF':
        ws[f'{col}{r}'].fill = fill(BLACK)
        ws[f'{col}{r}'].font = Font(name=F, size=9, bold=True, color=GOLD)
    if heads:
        for col, h in heads.items():
            ws[f'{col}{r}'] = h
            ws[f'{col}{r}'].alignment = Alignment(horizontal='center')
    ws.row_dimensions[r].height = 20
    ws[f'B{r}'].alignment = Alignment(vertical='center', indent=1)


def label(ws, r, text, note=None):
    ws[f'B{r}'] = text
    ws[f'B{r}'].font = Font(name=F, size=10, color=INK)
    ws[f'B{r}'].alignment = Alignment(indent=1, vertical='center', wrap_text=True)
    if note:
        ws[f'F{r}'] = note
        ws[f'F{r}'].font = Font(name=F, size=8, italic=True, color=GREY)
        ws[f'F{r}'].alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[r].height = 22


def inp(ws, ref, value=None, fmt=None):
    c = ws[ref]
    c.value = value
    c.fill = fill(CREAM)
    c.font = INPUT_FONT
    c.border = box
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.protection = Protection(locked=False)
    if fmt:
        c.number_format = fmt


def calc(ws, ref, formula, fmt=NUM, bold=False, gold=False):
    c = ws[ref]
    c.value = formula
    c.font = Font(name=F, size=11, bold=bold, color=BLACK if not gold else BLACK)
    c.number_format = fmt
    c.border = box
    c.alignment = Alignment(horizontal='center', vertical='center')
    if gold:
        c.fill = fill(GOLD)


def build_calc(ws, ex=None):
    ex = ex or {}
    ws.sheet_view.showGridLines = False
    for col, w in {'A': 2, 'B': 46, 'C': 17, 'D': 17, 'E': 17, 'F': 44, 'G': 2}.items():
        ws.column_dimensions[col].width = w
    title(ws, 'KAIZEN XP  |  KPI IMPACT CALCULATOR',
          'Put a pound sign on any problem. Fill in the cream boxes (blue text). Everything else works itself out.')

    section(ws, 5, 'THE ISSUE')
    for r, (lab, key) in enumerate([('What is the problem or improvement?', 'issue'), ('Line, machine or area', 'area'),
                                    ('Calculated by', 'by'), ('Date', 'date')], 6):
        label(ws, r, lab)
        ws.merge_cells(f'C{r}:E{r}')
        inp(ws, f'C{r}', ex.get(key))
        ws[f'C{r}'].alignment = Alignment(horizontal='left', vertical='center', indent=1)
        for col in 'DE':
            ws[f'{col}{r}'].fill = fill(CREAM); ws[f'{col}{r}'].border = box
    label(ws, 10, 'Main impact', 'Pick from the list.')
    inp(ws, 'C10', ex.get('impact'))
    dv = DataValidation(type='list', formula1='"Downtime,Labour time,Waste,Lost sales,Safety,Quality,More than one"', allow_blank=True)
    ws.add_data_validation(dv); dv.add('C10')

    section(ws, 12, '1.  HOW OFTEN AND HOW LONG', {'C': 'BEFORE', 'D': 'AFTER THE FIX'})
    rows = [(13, 'Times it happens per shift', 'times', 'Count it for a week if you can. A guess you can defend is fine to start.'),
            (14, 'Minutes lost each time', 'mins', 'From the moment it goes wrong to running normally again.'),
            (15, 'People tied up each time', 'people', 'How many people stop what they are doing to deal with it.')]
    for r, lab, key, note in rows:
        label(ws, r, lab, note)
        inp(ws, f'C{r}', ex.get(key + '_b'), NUM1)
        inp(ws, f'D{r}', ex.get(key + '_a'), NUM1)
    ws['F15'].value = 'How many people stop what they are doing. Leave AFTER blank if nothing changes.'
    label(ws, 16, 'Shifts per week it happens on', 'Example: 3 shifts a day, 5 days a week = 15.')
    inp(ws, 'C16', ex.get('shifts', 15), NUM)
    calc(ws, 'D16', '=C16')
    label(ws, 17, 'Weeks a year the job runs', 'Use 52, or fewer for shutdowns and closed weeks.')
    inp(ws, 'C17', ex.get('weeks', 52), NUM)
    calc(ws, 'D17', '=C17')

    section(ws, 19, '2.  WHAT AN HOUR COSTS')
    label(ws, 20, 'Staff cost per hour (£)', 'Not sure? Use your own hourly rate. It is the lowest honest number.')
    inp(ws, 'C20', ex.get('rate'), '£#,##0.00')
    label(ws, 21, 'Cost of an hour of line downtime (£)  OPTIONAL', 'Only if the problem stops a line. Ask what an hour of lost output is worth. Leave blank if not.')
    inp(ws, 'C21', ex.get('down'), '£#,##0')

    section(ws, 23, '3.  WASTE AND LOST SALES  (OPTIONAL)', {'C': 'BEFORE', 'D': 'AFTER THE FIX'})
    label(ws, 24, 'Units scrapped or wasted per week', 'Packs, kg, labels, products. Whatever ends up in the bin.')
    inp(ws, 'C24', ex.get('units_b'), NUM); inp(ws, 'D24', ex.get('units_a'), NUM)
    label(ws, 25, 'Cost per wasted unit (£)')
    inp(ws, 'C25', ex.get('unitcost'), '£#,##0.00')
    calc(ws, 'D25', '=C25', '£#,##0.00;-£#,##0.00;-')
    label(ws, 26, 'Sales lost per week (£)', 'Retail: empty shelves, out of stocks, customers walking out.')
    inp(ws, 'C26', ex.get('sales_b'), '£#,##0'); inp(ws, 'D26', ex.get('sales_a'), '£#,##0')
    label(ws, 27, 'Margin on those sales (%)', 'Type 25 for 25%. Only the margin is really lost, not the full price.')
    inp(ws, 'C27', ex.get('margin'), '0.0')
    calc(ws, 'D27', '=C27', '0.0;-0.0;-')

    section(ws, 29, 'WHAT IT COSTS A YEAR', {'C': 'BEFORE', 'D': 'AFTER THE FIX', 'E': 'SAVING'})
    # effective after values: blank AFTER means unchanged
    eff = lambda r: f'IF(D{r}="",N(C{r}),N(D{r}))'
    res = [
        (30, 'People hours lost a year', f'=N(C13)*N(C14)*N(C15)*N(C16)*N(C17)/60',
         f'={eff(13)}*{eff(14)}*{eff(15)}*N(D16)*N(D17)/60', NUM),
        (31, 'Line hours stopped a year', f'=IF(N(C21)=0,0,N(C13)*N(C14)*N(C16)*N(C17)/60)',
         f'=IF(N(C21)=0,0,{eff(13)}*{eff(14)}*N(D16)*N(D17)/60)', NUM),
        (32, 'Staff time cost', '=C30*N(C20)', '=D30*N(C20)', GBP),
        (33, 'Downtime cost', '=C31*N(C21)', '=D31*N(C21)', GBP),
        (34, 'Waste cost', '=N(C24)*N(C25)*N(C17)', f'={eff(24)}*N(C25)*N(C17)', GBP),
        (35, 'Lost margin', '=N(C26)*N(C27)/100*N(C17)', f'={eff(26)}*N(C27)/100*N(C17)', GBP),
    ]
    for r, lab, fb, fa, fmt in res:
        label(ws, r, lab)
        calc(ws, f'C{r}', fb, fmt)
        calc(ws, f'D{r}', fa, fmt)
        calc(ws, f'E{r}', f'=C{r}-D{r}', fmt, bold=True)
    ws['F31'] = 'Only counts if you filled in a downtime cost.'
    ws['F31'].font = Font(name=F, size=8, italic=True, color=GREY)
    label(ws, 36, 'TOTAL COST OF THIS PROBLEM A YEAR')
    ws['B36'].font = Font(name=F, size=11, bold=True, color=INK)
    calc(ws, 'C36', '=SUM(C32:C35)', GBP, bold=True)
    calc(ws, 'D36', '=SUM(D32:D35)', GBP, bold=True)
    calc(ws, 'E36', '=C36-D36', GBP, bold=True, gold=True)
    ws.row_dimensions[36].height = 26

    section(ws, 38, 'THE PAYBACK')
    label(ws, 39, 'What the fix cost (£)', 'Parts, materials, hours spent. Leave blank if it cost nothing.')
    inp(ws, 'C39', ex.get('fixcost'), '£#,##0')
    label(ws, 40, 'Saving per week (£)')
    calc(ws, 'C40', '=IF(N(C17)=0,0,E36/C17)', GBP)
    label(ws, 41, 'The fix pays for itself in')
    calc(ws, 'C41', '=IF(E36<=0,"-",IF(N(C39)=0,"Straight away",IF(C39/C40<1,"Under a week",ROUND(C39/C40,0)&" weeks")))', 'General', bold=True)
    ws.merge_cells('C41:D41')

    section(ws, 43, 'YOUR SENTENCE   (copy it into your Weekly Wins email, improvement log and review)')
    ws.merge_cells('B44:F46')
    ws['B44'] = ('=IF(N(C13)=0,"Fill in the cream boxes and your sentence appears here.",'
                 '"I reduced "&IF(C6="","the problem",C6)&" from "&C13&" to "&' + eff(13) + '&" times a shift. '
                 'That gives back about "&TEXT(E30,"#,##0")&" people hours a year and is worth roughly £"&TEXT(E36,"#,##0")&" a year'
                 '"&IF(N(C39)>0,", for a fix that cost £"&TEXT(C39,"#,##0"),"")&".")')
    ws['B44'].font = Font(name=F, size=12, italic=True, color=INK)
    ws['B44'].alignment = Alignment(wrap_text=True, vertical='center', indent=1)
    ws['B44'].fill = fill(CREAM)
    for r in (44, 45, 46):
        ws.row_dimensions[r].height = 22
    ws['B48'] = '© 2026 Kaizen XP. Licensed for one person\'s use. Please do not share or resell.'
    ws['B48'].font = Font(name=F, size=8, color='999999')
    ws.freeze_panes = 'A4'
    ws.print_area = 'A1:G48'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


# ------------------------------------------------ START HERE
st = wb.active
st.title = 'Start Here'
st.sheet_view.showGridLines = False
st.column_dimensions['A'].width = 2
st.column_dimensions['B'].width = 26
st.column_dimensions['C'].width = 80
title(st, 'KAIZEN XP  |  KPI IMPACT CALCULATOR', 'Put a pound sign on any problem in two minutes.', last='C')
rows = [
    ('WHAT IT DOES', 'Turns "this keeps happening" into a number a manager will act on: hours lost, what they cost, and what fixing it is worth every year.'),
    ('WHY IT MATTERS', '"The labeller keeps jamming" gets a nod. "The labeller jams are costing us £4,300 a year and the fix costs £25" gets a decision.'),
    ('', ''),
    ('HOW TO USE IT', '1. Go to the Calculator tab.'),
    ('', '2. Fill in the cream boxes with blue text. Leave anything that does not apply blank.'),
    ('', '3. Fill in BEFORE now. Fill in AFTER THE FIX once your fix has been in for a week or two.'),
    ('', '4. Copy your sentence from the bottom of the page into your email or review.'),
    ('', '5. Log every result on the Improvement Log tab so you have a running total.'),
    ('', ''),
    ('THE TABS', 'Example: a filled-in calculation (Line 3 labeller) so you can see what good looks like.'),
    ('', 'Calculator: your own calculation. Save a copy of the file for each new problem.'),
    ('', 'Improvement Log: every problem you have costed and fixed, with a running total.'),
    ('', ''),
    ('THE RULES', 'Count, do not guess, where you can. A week of tally marks beats a hunch.'),
    ('', 'If you do not know what an hour costs, use your own hourly rate. It is the lowest honest number, so nobody can argue with it.'),
    ('', 'Only count margin on lost sales, not the full price. It keeps your numbers believable.'),
    ('', ''),
    ('WORKS IN', 'Microsoft Excel, Google Sheets (File > Import) and Apple Numbers. The formulas are locked so they cannot be broken by accident. The cream boxes are not.'),
]
r = 5
for a, b in rows:
    st[f'B{r}'] = a
    st[f'B{r}'].font = Font(name=F, size=9, bold=True, color=INK)
    st[f'C{r}'] = b
    st[f'C{r}'].font = Font(name=F, size=10, color=INK)
    st[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    st[f'B{r}'].alignment = Alignment(vertical='top')
    st.row_dimensions[r].height = 30 if len(b) > 85 else 17
    r += 1
r += 1
st[f'B{r}'] = 'LEGEND'
st[f'B{r}'].font = Font(name=F, size=9, bold=True, color=INK)
st[f'C{r}'] = 'Type here'
st[f'C{r}'].fill = fill(CREAM); st[f'C{r}'].font = INPUT_FONT; st[f'C{r}'].border = box
r += 1
st[f'C{r}'] = 'Worked out for you (locked)'
st[f'C{r}'].font = Font(name=F, size=10, color=BLACK); st[f'C{r}'].border = box
r += 1
st[f'C{r}'] = 'The headline number'
st[f'C{r}'].fill = fill(GOLD); st[f'C{r}'].font = Font(name=F, size=10, bold=True, color=BLACK); st[f'C{r}'].border = box
r += 2
st[f'B{r}'] = '"Do what nobody else can be bothered to do, and you\'ll get what nobody else gets."'
st[f'B{r}'].font = Font(name=F, size=10, italic=True, color=GREY)
st[f'B{r + 1}'] = '© 2026 Kaizen XP. Licensed for one person\'s use. Please do not share or resell.'
st[f'B{r + 1}'].font = Font(name=F, size=8, color='999999')

st.page_setup.fitToWidth = 1
st.page_setup.fitToHeight = 1
st.sheet_properties.pageSetUpPr.fitToPage = True
# ------------------------------------------------ EXAMPLE + CALCULATOR
exs = wb.create_sheet('Example')
build_calc(exs, ex={'issue': 'Line 3 labeller jams at reel change', 'area': 'Line 3, nights', 'by': 'Sam', 'date': 'Week 2',
                    'impact': 'More than one', 'times_b': 4, 'times_a': 1, 'mins_b': 6, 'mins_a': None, 'people_b': 1, 'people_a': None,
                    'shifts': 15, 'weeks': 52, 'rate': 13, 'down': None, 'units_b': 40, 'units_a': 10, 'unitcost': 0.8,
                    'sales_b': None, 'sales_a': None, 'margin': None, 'fixcost': 25})
exs['F6'] = 'This tab is a worked example. Use the Calculator tab for your own.'
exs['F6'].font = Font(name=F, size=8, bold=True, italic=True, color='9A7B1F')
exs['F6'].alignment = Alignment(wrap_text=True, vertical='center')
exs.sheet_properties.tabColor = GOLD

calc_ws = wb.create_sheet('Calculator')
build_calc(calc_ws)
calc_ws.sheet_properties.tabColor = BLACK

# ------------------------------------------------ IMPROVEMENT LOG
lg = wb.create_sheet('Improvement Log')
lg.sheet_view.showGridLines = False
widths = {'A': 2, 'B': 12, 'C': 40, 'D': 20, 'E': 16, 'F': 16, 'G': 22, 'H': 18}
for col, w in widths.items():
    lg.column_dimensions[col].width = w
lg.merge_cells('B1:H1'); lg['B1'] = 'KAIZEN XP  |  IMPROVEMENT LOG'
lg['B1'].font = Font(name=F, size=16, bold=True, color=GOLD); lg['B1'].alignment = Alignment(vertical='center')
lg.merge_cells('B2:H2'); lg['B2'] = 'Every problem you have costed and fixed. Copy the SAVING numbers across from each calculation.'
lg['B2'].font = Font(name=F, size=9, italic=True, color='BBBBBB')
for col in 'ABCDEFGHI':
    lg[f'{col}1'].fill = fill(BLACK); lg[f'{col}2'].fill = fill(BLACK); lg[f'{col}3'].fill = fill(GOLD)
lg.row_dimensions[1].height = 34; lg.row_dimensions[3].height = 3
# totals
lg['B5'] = 'RUNNING TOTAL'
lg['B5'].font = Font(name=F, size=9, bold=True, color=INK)
lg['D5'] = 'Improvements logged'; lg['D5'].font = Font(name=F, size=9, color=GREY)
lg['E5'] = '=COUNTA(C9:C58)'
lg['F5'] = '=SUM(F9:F58)'
lg['E6'] = '=SUM(E9:E58)'
lg['D6'] = 'Hours a year'; lg['D6'].font = Font(name=F, size=9, color=GREY)
lg['G5'] = '£ a year saved, all fixes'; lg['G5'].font = Font(name=F, size=9, color=GREY)
for ref, fmt in [('E5', NUM), ('F5', GBP), ('E6', NUM)]:
    lg[ref].font = Font(name=F, size=12, bold=True, color=BLACK); lg[ref].number_format = fmt
    lg[ref].alignment = Alignment(horizontal='center'); lg[ref].border = box
lg['F5'].fill = fill(GOLD)
heads = ['DATE', 'PROBLEM FIXED', 'LINE / AREA', 'HOURS SAVED A YEAR', '£ SAVED A YEAR', 'WHO SAW IT?', 'STATUS']
for i, h in enumerate(heads):
    cell = lg.cell(row=8, column=2 + i, value=h)
    cell.fill = fill(BLACK); cell.font = Font(name=F, size=9, bold=True, color=GOLD)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
lg.row_dimensions[8].height = 28
example = ['Week 2', 'Line 3 labeller jams at reel change', 'Line 3', 234, 4290, 'Line leader and manager', 'Fixed and checked']
for r in range(9, 59):
    for i in range(7):
        cell = lg.cell(row=r, column=2 + i)
        cell.border = box
        cell.font = INPUT_FONT
        cell.alignment = Alignment(horizontal='left' if i in (1, 2) else 'center', vertical='center')
        if i == 3: cell.number_format = NUM
        if i == 4: cell.number_format = GBP
    lg.row_dimensions[r].height = 18
for i, v in enumerate(example):
    c = lg.cell(row=9, column=2 + i, value=v)
    c.fill = fill(CREAM); c.font = Font(name=F, size=10, italic=True, color='0000FF')
lg['I9'] = 'Example row. Overwrite it.'
lg['I9'].font = Font(name=F, size=8, italic=True, color=GREY)
dv1 = DataValidation(type='list', formula1='"Line leader,Manager,Line leader and manager,Wider team,Nobody yet"', allow_blank=True)
dv2 = DataValidation(type='list', formula1='"Costed only,Fix in progress,Fixed,Fixed and checked"', allow_blank=True)
lg.add_data_validation(dv1); lg.add_data_validation(dv2)
dv1.add('G9:G58'); dv2.add('H9:H58')
lg.freeze_panes = 'B9'

wb.active = 0
wb.properties.title = 'KPI Impact Calculator'
wb.properties.creator = 'Kaizen XP'
wb.save(OUT)
print('saved')
