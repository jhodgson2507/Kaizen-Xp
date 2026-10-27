import sys, random, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

OUT = sys.argv[1]
BLACK, GOLD, CREAM, INK, GREY, LINE = '0D0D0D', 'C9A84C', 'F3EFE4', '1A1A1A', '666666', 'D3CCB8'
F = 'Arial'
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)
thin = Side(style='thin', color=LINE)
box = Border(left=thin, right=thin, top=thin, bottom=thin)
INPUT = Font(name=F, size=10, color='0000FF')
NUM = '#,##0;(#,##0);-'
GBP = '£#,##0;(£#,##0);-'
FIRST, LAST = 7, 506          # log rows
MACH_ROWS = list(range(15, 30))  # Setup machine list B15:B29
TYPE_ROWS = list(range(15, 25))  # Setup failure type list D15:D24
SHIFTS = ['Days', 'Earlies', 'Lates', 'Nights', 'Weekend']

wb = Workbook()


def bar(ws, cols, rows=(1, 2, 3)):
    for col in cols:
        for r in rows[:2]:
            ws[f'{col}{r}'].fill = fill(BLACK)
        ws[f'{col}{rows[2]}'].fill = fill(GOLD)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 3


def head(ws, title, sub, last):
    cols = [chr(c) for c in range(ord('A'), ord(last) + 2)]
    bar(ws, cols)
    ws.merge_cells(f'B1:{last}1'); ws['B1'] = title
    ws['B1'].font = Font(name=F, size=16, bold=True, color=GOLD); ws['B1'].alignment = Alignment(vertical='center')
    ws.merge_cells(f'B2:{last}2'); ws['B2'] = sub
    ws['B2'].font = Font(name=F, size=9, italic=True, color='BBBBBB'); ws['B2'].alignment = Alignment(vertical='top')
    ws.sheet_view.showGridLines = False


def sect(ws, r, c1, c2, text):
    for c in range(ord(c1), ord(c2) + 1):
        cell = ws[f'{chr(c)}{r}']
        cell.fill = fill(BLACK); cell.font = Font(name=F, size=9, bold=True, color=GOLD)
    ws[f'{c1}{r}'] = text
    ws[f'{c1}{r}'].alignment = Alignment(indent=1, vertical='center')
    ws.row_dimensions[r].height = 20


def inp(cell, v=None, fmt=None):
    cell.value = v; cell.fill = fill(CREAM); cell.font = INPUT; cell.border = box
    cell.protection = Protection(locked=False)
    cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    if fmt: cell.number_format = fmt


# ======================================================= START HERE
st = wb.active; st.title = 'Start Here'
st.column_dimensions['A'].width = 2; st.column_dimensions['B'].width = 24; st.column_dimensions['C'].width = 84
head(st, 'KAIZEN XP  |  BREAKDOWN ANALYSIS DASHBOARD', 'Log every breakdown. See which machine to fix first, and prove it got better.', 'C')
rows = [
    ('WHAT IT DOES', 'Turns a simple breakdown log into MTBF, MTTR, availability, downtime cost and a Pareto of your worst machines and failure types. Updates itself every time you add a line.'),
    ('WHY IT MATTERS', 'Every site has a breakdown log. Almost nobody turns it into a decision. This does it for you: the dashboard tells you which machine and which failure type to attack first.'),
    ('', ''),
    ('HOW TO USE IT', '1. Look at the Dashboard first. It is running on SAMPLE DATA so you can see it working.'),
    ('', '2. When you are ready, go to Breakdown Log, select the sample rows and press Delete.'),
    ('', '3. On Setup, type your line name, the period you want to look at, your run hours and your machine list.'),
    ('', '4. Log every breakdown: date, shift, machine, failure type, minutes down. Pick from the lists.'),
    ('', '5. Check the Dashboard weekly. Attack the top bar on the Pareto with 5 Whys or an RCFA.'),
    ('', ''),
    ('THE NUMBERS', 'MTBF (Mean Time Between Failures): run hours divided by breakdowns. Higher is better.'),
    ('', 'MTTR (Mean Time To Repair): downtime minutes divided by breakdowns. Lower is better.'),
    ('', 'Availability: the share of planned run time the line was not broken down.'),
    ('', 'Pareto: machines and failure types in order of downtime, with a running total. The first one or two bars are usually most of the problem.'),
    ('', ''),
    ('BREAKDOWN OR ASSIST?', 'A breakdown needs a technician or engineer and stops the job. The quick stops an operator clears in a minute are assists: log those in the Kaizen XP MTBA Assist Tracker.'),
    ('', ''),
    ('WORKS IN', 'Microsoft Excel, Google Sheets (File > Import) and Apple Numbers. The Dashboard and Workings are locked so formulas cannot be broken by accident.'),
]
r = 5
for a, b in rows:
    st[f'B{r}'] = a; st[f'B{r}'].font = Font(name=F, size=9, bold=True, color=INK); st[f'B{r}'].alignment = Alignment(vertical='top')
    st[f'C{r}'] = b; st[f'C{r}'].font = Font(name=F, size=10, color=INK); st[f'C{r}'].alignment = Alignment(wrap_text=True, vertical='top')
    st.row_dimensions[r].height = 30 if len(b) > 90 else 17
    r += 1
r += 1
st[f'B{r}'] = 'LEGEND'; st[f'B{r}'].font = Font(name=F, size=9, bold=True, color=INK)
st[f'C{r}'] = 'Type here'; st[f'C{r}'].fill = fill(CREAM); st[f'C{r}'].font = INPUT; st[f'C{r}'].border = box
st[f'C{r + 1}'] = 'Worked out for you (locked)'; st[f'C{r + 1}'].border = box; st[f'C{r + 1}'].font = Font(name=F, size=10)
r += 3
st[f'B{r}'] = '"Do what nobody else can be bothered to do, and you\'ll get what nobody else gets."'
st[f'B{r}'].font = Font(name=F, size=10, italic=True, color=GREY)
st[f'B{r + 1}'] = '© 2026 Kaizen XP. Licensed for one person\'s use. Please do not share or resell.'
st[f'B{r + 1}'].font = Font(name=F, size=8, color='999999')

# ======================================================= DASHBOARD (create early for tab order)
db = wb.create_sheet('Dashboard')
log = wb.create_sheet('Breakdown Log')
su = wb.create_sheet('Setup')
wk = wb.create_sheet('Workings')

# ======================================================= SETUP
su.column_dimensions['A'].width = 2
for col, w in {'B': 34, 'C': 22, 'D': 30, 'E': 4, 'F': 50}.items():
    su.column_dimensions[col].width = w
head(su, 'SETUP', 'Tell the dashboard about your line. Cream boxes only. The sample values are there to show you the format.', 'F')
sect(su, 5, 'B', 'D', 'YOUR LINE AND PERIOD')
items = [(6, 'Line or area name', 'Line 3 (sample)', None, ''),
         (7, 'Period start date', dt.date(2026, 4, 1), 'dd mmm yyyy', 'The dashboard only counts breakdowns between these dates.'),
         (8, 'Period end date', dt.date(2026, 9, 30), 'dd mmm yyyy', 'Change both to look at a month, a quarter or a year.'),
         (9, 'Planned run hours per day', 21, '0.0', 'Hours the line is meant to run, not counting planned stops.'),
         (10, 'Run days per week', 5, '0', ''),
         (11, 'Cost of an hour of downtime (£)', 250, '£#,##0', 'Ask what an hour of lost output is worth. If unsure, use labour cost as the floor.')]
for rr, lab, v, fmt, note in items:
    su[f'B{rr}'] = lab; su[f'B{rr}'].font = Font(name=F, size=10, color=INK); su[f'B{rr}'].alignment = Alignment(indent=1)
    inp(su[f'C{rr}'], v, fmt)
    if note:
        su[f'D{rr}'] = note; su[f'D{rr}'].font = Font(name=F, size=8, italic=True, color=GREY)
        su[f'D{rr}'].alignment = Alignment(wrap_text=True, vertical='center')
    su.row_dimensions[rr].height = 22
su.merge_cells('D9:F9'); su.merge_cells('D11:F11'); su.merge_cells('D7:F7'); su.merge_cells('D8:F8')
su['B12'] = 'Planned run hours in this period'; su['B12'].font = Font(name=F, size=10, bold=True, color=INK); su['B12'].alignment = Alignment(indent=1)
su['C12'] = '=IF(OR(C7="",C8=""),0,(C8-C7+1)/7*N(C10)*N(C9))'
su['C12'].number_format = '#,##0'; su['C12'].font = Font(name=F, size=10, bold=True); su['C12'].border = box
sect(su, 14, 'B', 'B', 'MACHINES (UP TO 15)')
sect(su, 14, 'D', 'D', 'FAILURE TYPES (UP TO 10)')
machines = ['Depalletiser', 'Filler', 'Capper', 'Labeller', 'Shrink wrapper', 'Case packer', 'Palletiser', 'Conveyors']
types = ['Mechanical', 'Electrical', 'Pneumatic', 'Sensor or controls', 'Changeover or setup', 'Material or supply', 'Other']
for i, rr in enumerate(MACH_ROWS):
    inp(su[f'B{rr}'], machines[i] if i < len(machines) else None)
for i, rr in enumerate(TYPE_ROWS):
    inp(su[f'D{rr}'], types[i] if i < len(types) else None)
su['F15'] = 'These lists feed the dropdowns in the Breakdown Log and the charts on the Dashboard. Rename them to match your site.'
su['F15'].font = Font(name=F, size=8, italic=True, color=GREY); su['F15'].alignment = Alignment(wrap_text=True, vertical='top')
su.merge_cells('F15:F18')

# ======================================================= LOG
log.column_dimensions['A'].width = 2
widths = {'B': 13, 'C': 11, 'D': 18, 'E': 20, 'F': 46, 'G': 12, 'H': 14, 'I': 15, 'J': 12}
for col, w in widths.items():
    log.column_dimensions[col].width = w
head(log, 'BREAKDOWN LOG', 'One line per breakdown. Pick machine, failure type and shift from the lists. Minutes down = from stopped to running again.', 'J')
log.merge_cells('B4:J4')
log['B4'] = ('The rows below are SAMPLE DATA so you can see the dashboard working. When you are ready: select them, press Delete, '
             'then set your own machines and dates on the Setup tab.')
log['B4'].font = Font(name=F, size=9, bold=True, italic=True, color='9A7B1F'); log['B4'].alignment = Alignment(wrap_text=True, vertical='center')
log.row_dimensions[4].height = 30
heads = ['DATE', 'SHIFT', 'MACHINE', 'FAILURE TYPE', 'WHAT HAPPENED', 'MINUTES DOWN', 'FIXED BY', 'ROOT CAUSE FOUND?', 'RCFA REF']
for i, h in enumerate(heads):
    c = log.cell(row=6, column=2 + i, value=h)
    c.fill = fill(BLACK); c.font = Font(name=F, size=9, bold=True, color=GOLD)
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
log.row_dimensions[6].height = 28
for r in range(FIRST, LAST + 1):
    for i in range(9):
        c = log.cell(row=r, column=2 + i)
        c.border = box; c.font = INPUT
        c.alignment = Alignment(horizontal='left' if i in (2, 3, 4, 6) else 'center', vertical='center')
    log.cell(row=r, column=2).number_format = 'dd mmm yy'
    log.cell(row=r, column=7).number_format = '0'
dvs = [DataValidation(type='list', formula1='"' + ','.join(SHIFTS) + '"', allow_blank=True),
       DataValidation(type='list', formula1=f'=Setup!$B${MACH_ROWS[0]}:$B${MACH_ROWS[-1]}', allow_blank=True),
       DataValidation(type='list', formula1=f'=Setup!$D${TYPE_ROWS[0]}:$D${TYPE_ROWS[-1]}', allow_blank=True),
       DataValidation(type='list', formula1='"Yes,No,In progress"', allow_blank=True),
       DataValidation(type='whole', operator='between', formula1='0', formula2='10000', allow_blank=True,
                      errorTitle='Minutes down', error='Type the minutes as a whole number, for example 45.')]
for dv in dvs:
    log.add_data_validation(dv)
dvs[0].add(f'C{FIRST}:C{LAST}'); dvs[1].add(f'D{FIRST}:D{LAST}'); dvs[2].add(f'E{FIRST}:E{LAST}')
dvs[3].add(f'I{FIRST}:I{LAST}'); dvs[4].add(f'G{FIRST}:G{LAST}')
log.freeze_panes = f'B{FIRST}'

# sample data
random.seed(7)
desc = {
    'Mechanical': ['Bearing noise, replaced', 'Drive chain off sprocket', 'Guide rail worked loose', 'Gearbox coupling sheared'],
    'Electrical': ['Motor tripped on overload', 'Contactor failed', 'Loose terminal in panel'],
    'Pneumatic': ['Air leak on cylinder', 'Valve stuck, replaced', 'Low air pressure alarm'],
    'Sensor or controls': ['Photo eye misaligned', 'Proximity sensor failed', 'PLC fault, reset and checked'],
    'Changeover or setup': ['Wrong change parts fitted', 'Settings not reset after changeover', 'Overran changeover'],
    'Material or supply': ['Label reel splice failed', 'Out of spec cases jammed', 'Film roll faulty'],
    'Other': ['Guard interlock fault', 'Lubrication missed, seized'],
}
labeller_types = (['Sensor or controls'] * 3 + ['Material or supply'] * 4 + ['Mechanical'] * 2 + ['Changeover or setup'])
mw = {'Labeller': 34, 'Filler': 17, 'Case packer': 14, 'Capper': 10, 'Shrink wrapper': 10, 'Palletiser': 8, 'Conveyors': 7, 'Depalletiser': 4}
tw = {'Mechanical': 26, 'Electrical': 14, 'Pneumatic': 14, 'Sensor or controls': 16, 'Changeover or setup': 12, 'Material or supply': 12, 'Other': 6}
fixers = ['Day tech', 'Night tech', 'Electrician', 'Shift engineer', 'Operator and tech']
d = dt.date(2026, 4, 1)
recs = []
while d <= dt.date(2026, 9, 30):
    if d.weekday() < 5:
        for _ in range(random.choice([0, 1, 1, 1, 2])):
            after_fix = d >= dt.date(2026, 8, 1)
            weights = dict(mw)
            if after_fix:
                weights['Labeller'] = 8
            m = random.choices(list(weights), list(weights.values()))[0]
            t = random.choice(labeller_types) if m == 'Labeller' else random.choices(list(tw), list(tw.values()))[0]
            mins = int(min(240, max(8, random.lognormvariate(3.3, 0.6))))
            if m == 'Labeller' and not after_fix:
                mins = int(mins * 1.6)
            rc = random.choices(['Yes', 'No', 'In progress'], [40, 35, 25])[0]
            recs.append([d, random.choice(['Days', 'Lates', 'Nights', 'Earlies']), m, t, random.choice(desc[t]), mins,
                         random.choice(fixers), rc, ''])
    d += dt.timedelta(days=1)
n_rc = 0
for i, rec in enumerate(recs):
    if rec[7] == 'Yes' and random.random() < 0.5:
        n_rc += 1
        rec[8] = f'RCFA-{n_rc:03d}'
    for j, v in enumerate(rec):
        log.cell(row=FIRST + i, column=2 + j, value=v)
print('sample rows', len(recs), 'last row', FIRST + len(recs) - 1)

# ======================================================= WORKINGS
wk.column_dimensions['A'].width = 2
for col, w in {'B': 24, 'C': 12, 'D': 12, 'E': 14, 'F': 4, 'G': 8, 'H': 24, 'I': 12, 'J': 12, 'K': 12, 'L': 12}.items():
    wk.column_dimensions[col].width = w
head(wk, 'WORKINGS', 'The maths behind the Dashboard. Nothing to type here.', 'L')
R = lambda col: f"'Breakdown Log'!${col}${FIRST}:${col}${LAST}"
DATES, SHIFT, MACH, TYPE, MINS, RC = R('B'), R('C'), R('D'), R('E'), R('G'), R('I')
PER = f'{DATES},">="&Setup!$C$7,{DATES},"<="&Setup!$C$8'

wk['B5'] = 'BY MACHINE (setup order)'; wk['H5'] = 'BY MACHINE (worst first)'
for c, h in zip('BCDE', ['Machine', 'Breakdowns', 'Minutes', 'Sort key']):
    wk[f'{c}6'] = h
for c, h in zip('GHIJKL', ['Rank', 'Machine', 'Minutes', 'Hours', 'Running %', 'Breakdowns']):
    wk[f'{c}6'] = h
for i, srow in enumerate(MACH_ROWS):
    r = 7 + i
    wk[f'B{r}'] = f'=IF(Setup!$B${srow}="","",Setup!$B${srow})'
    wk[f'C{r}'] = f'=IF(B{r}="",0,COUNTIFS({MACH},B{r},{PER}))'
    wk[f'D{r}'] = f'=IF(B{r}="",0,SUMIFS({MINS},{MACH},B{r},{PER}))'
    wk[f'E{r}'] = f'=IF(B{r}="",-1,D{r}+(100-ROW())/1000)'
    k = i + 1
    wk[f'G{r}'] = k
    key = f'LARGE($E$7:$E$21,G{r})'
    wk[f'H{r}'] = f'=IF({key}<0,"",INDEX($B$7:$B$21,MATCH({key},$E$7:$E$21,0)))'
    wk[f'I{r}'] = f'=IF(H{r}="",0,INDEX($D$7:$D$21,MATCH({key},$E$7:$E$21,0)))'
    wk[f'J{r}'] = f'=I{r}/60'
    wk[f'K{r}'] = f'=IF(SUM($I$7:$I$21)=0,0,SUM($I$7:I{r})/SUM($I$7:$I$21))'
    wk[f'L{r}'] = f'=IF(H{r}="",0,INDEX($C$7:$C$21,MATCH({key},$E$7:$E$21,0)))'
    wk[f'J{r}'].number_format = '0.0'; wk[f'K{r}'].number_format = '0%'

wk['B24'] = 'BY FAILURE TYPE (setup order)'; wk['H24'] = 'BY FAILURE TYPE (worst first)'
for c, h in zip('BCDE', ['Failure type', 'Breakdowns', 'Minutes', 'Sort key']):
    wk[f'{c}25'] = h
for c, h in zip('GHIJKL', ['Rank', 'Failure type', 'Minutes', 'Hours', 'Running %', 'Breakdowns']):
    wk[f'{c}25'] = h
for i, srow in enumerate(TYPE_ROWS):
    r = 26 + i
    wk[f'B{r}'] = f'=IF(Setup!$D${srow}="","",Setup!$D${srow})'
    wk[f'C{r}'] = f'=IF(B{r}="",0,COUNTIFS({TYPE},B{r},{PER}))'
    wk[f'D{r}'] = f'=IF(B{r}="",0,SUMIFS({MINS},{TYPE},B{r},{PER}))'
    wk[f'E{r}'] = f'=IF(B{r}="",-1,D{r}+(100-ROW())/1000)'
    wk[f'G{r}'] = i + 1
    key = f'LARGE($E$26:$E$35,G{r})'
    wk[f'H{r}'] = f'=IF({key}<0,"",INDEX($B$26:$B$35,MATCH({key},$E$26:$E$35,0)))'
    wk[f'I{r}'] = f'=IF(H{r}="",0,INDEX($D$26:$D$35,MATCH({key},$E$26:$E$35,0)))'
    wk[f'J{r}'] = f'=I{r}/60'
    wk[f'K{r}'] = f'=IF(SUM($I$26:$I$35)=0,0,SUM($I$26:I{r})/SUM($I$26:$I$35))'
    wk[f'L{r}'] = f'=IF(H{r}="",0,INDEX($C$26:$C$35,MATCH({key},$E$26:$E$35,0)))'
    wk[f'J{r}'].number_format = '0.0'; wk[f'K{r}'].number_format = '0%'

wk['B38'] = 'BY MONTH'
for c, h in zip('BCDEG', ['Month start', 'Breakdowns', 'Hours down', 'Label', '']):
    if h: wk[f'{c}39'] = h
for i in range(12):
    r = 40 + i
    wk[f'B{r}'] = '=DATE(YEAR(Setup!$C$7),MONTH(Setup!$C$7),1)' if i == 0 else f'=EDATE(B{r - 1},1)'
    wk[f'B{r}'].number_format = 'dd mmm yyyy'
    nxt = f'EDATE(B{r},1)'
    wk[f'C{r}'] = f'=COUNTIFS({DATES},">="&B{r},{DATES},"<"&{nxt},{PER})'
    wk[f'D{r}'] = f'=SUMIFS({MINS},{DATES},">="&B{r},{DATES},"<"&{nxt},{PER})/60'
    wk[f'D{r}'].number_format = '0.0'
    wk[f'E{r}'] = f'=TEXT(B{r},"mmm yy")'

wk['H38'] = 'BY SHIFT'
for c, h in zip('HIJ', ['Shift', 'Breakdowns', 'Hours down']):
    wk[f'{c}39'] = h
for i, s in enumerate(SHIFTS):
    r = 40 + i
    wk[f'H{r}'] = s
    wk[f'I{r}'] = f'=COUNTIFS({SHIFT},H{r},{PER})'
    wk[f'J{r}'] = f'=SUMIFS({MINS},{SHIFT},H{r},{PER})/60'
    wk[f'J{r}'].number_format = '0.0'

wk['B55'] = 'TOTALS'
tot = [('Breakdowns', f'=COUNTIFS({PER})', '#,##0'),
       ('Minutes down', f'=SUMIFS({MINS},{PER})', '#,##0'),
       ('Planned run hours', '=Setup!$C$12', '#,##0'),
       ('MTBF (hours)', '=IF(C56=0,"-",(C58-C57/60)/C56)', '#,##0.0'),
       ('MTTR (minutes)', '=IF(C56=0,"-",C57/C56)', '#,##0'),
       ('Availability', '=IF(C58=0,"-",(C58-C57/60)/C58)', '0.0%'),
       ('Downtime cost', '=C57/60*N(Setup!$C$11)', GBP),
       ('Root cause found', f'=IF(C56=0,"-",COUNTIFS({RC},"Yes",{PER})/C56)', '0%'),
       ('Top machine share', '=IF(C57=0,0,I7/C57)', '0%'),
       ('Top failure share', '=IF(C57=0,0,I26/C57)', '0%')]
for i, (lab, f_, fmt) in enumerate(tot):
    r = 56 + i
    wk[f'B{r}'] = lab; wk[f'C{r}'] = f_; wk[f'C{r}'].number_format = fmt
for row in wk.iter_rows(min_row=5, max_row=66):
    for c in row:
        if c.value is not None and c.font.color is None or (c.value is not None and c.row > 4):
            c.font = Font(name=F, size=9, bold=c.row in (5, 6, 24, 25, 38, 39, 55), color=INK)
wk.protection.sheet = True
wk.sheet_properties.tabColor = '999999'

# ======================================================= DASHBOARD
db.sheet_view.showGridLines = False
db.column_dimensions['A'].width = 2
for col in 'BCDEFGHIJKLMNO':
    db.column_dimensions[col].width = 10.5
head(db, 'KAIZEN XP  |  BREAKDOWN ANALYSIS DASHBOARD', '="Line: "&Setup!C6&"     Period: "&TEXT(Setup!C7,"d mmm yyyy")&" to "&TEXT(Setup!C8,"d mmm yyyy")', 'O')
tiles = [('BREAKDOWNS', '=Workings!C56', '#,##0', 'B', 'C'), ('HOURS DOWN', '=Workings!C57/60', '#,##0.0', 'D', 'E'),
         ('MTBF (HOURS)', '=Workings!C59', '#,##0.0', 'F', 'G'), ('MTTR (MINS)', '=Workings!C60', '#,##0', 'H', 'I'),
         ('AVAILABILITY', '=Workings!C61', '0.0%', 'J', 'K'), ('DOWNTIME COST', '=Workings!C62', GBP, 'L', 'M'),
         ('ROOT CAUSE FOUND', '=Workings!C63', '0%', 'N', 'O')]
for lab, f_, fmt, a, b in tiles:
    db.merge_cells(f'{a}5:{b}5'); db.merge_cells(f'{a}6:{b}7')
    db[f'{a}5'] = lab
    db[f'{a}5'].font = Font(name=F, size=8, bold=True, color=GOLD)
    db[f'{a}5'].alignment = Alignment(horizontal='center', vertical='center')
    db[f'{a}6'] = f_
    db[f'{a}6'].number_format = fmt
    db[f'{a}6'].font = Font(name=F, size=20, bold=True, color='FFFFFF')
    db[f'{a}6'].alignment = Alignment(horizontal='center', vertical='center')
    for col in (a, b):
        for rr in (5, 6, 7):
            db[f'{col}{rr}'].fill = fill(BLACK)
db.row_dimensions[5].height = 20; db.row_dimensions[6].height = 22; db.row_dimensions[7].height = 22
for col in 'BCDEFGHIJKLMNO':
    db[f'{col}8'].fill = fill(GOLD)
db.row_dimensions[8].height = 3
db.merge_cells('B10:O11')
db['B10'] = ('=IF(Workings!C56=0,"Log your first breakdown and the machine to fix first shows here.",'
             '"FIX FIRST:  "&Workings!H7&". It caused "&TEXT(Workings!C64,"0%")&" of all downtime ("&TEXT(Workings!J7,"0.0")&" hours) in this period. '
             'Biggest failure type overall: "&Workings!H26&" ("&TEXT(Workings!C65,"0%")&").")')
db['B10'].font = Font(name=F, size=12, bold=True, color=INK)
db['B10'].alignment = Alignment(wrap_text=True, vertical='center', indent=1)
for col in 'BCDEFGHIJKLMNO':
    for rr in (10, 11):
        db[f'{col}{rr}'].fill = fill(CREAM)
db.row_dimensions[10].height = 20; db.row_dimensions[11].height = 20


def style_bar(ch, title, ytitle=None):
    ch.title = title
    ch.style = 2
    ch.height = 7.5
    ch.width = 14.6
    ch.legend = None
    if ytitle:
        ch.y_axis.title = ytitle
    ch.y_axis.majorGridlines = None
    return ch


# Pareto by machine
b1 = BarChart(); b1.type = 'col'
b1.add_data(Reference(wk, min_col=10, min_row=6, max_row=16), titles_from_data=True)
b1.set_categories(Reference(wk, min_col=8, min_row=7, max_row=16))
style_bar(b1, 'Downtime by machine, worst first (hrs)', 'Hours')
b1.series[0].graphicalProperties.solidFill = GOLD
b1.series[0].graphicalProperties.line.solidFill = GOLD
l1 = LineChart()
l1.add_data(Reference(wk, min_col=11, min_row=6, max_row=16), titles_from_data=True)
l1.y_axis.axId = 200; l1.y_axis.title = 'Running total'; l1.y_axis.number_format = '0%'
l1.y_axis.scaling.min = 0; l1.y_axis.scaling.max = 1; l1.y_axis.majorGridlines = None
l1.y_axis.crosses = 'max'
l1.series[0].graphicalProperties.line.solidFill = BLACK
l1.series[0].graphicalProperties.line.width = 22000
l1.series[0].smooth = False
b1.x_axis.delete = False; b1.y_axis.delete = False; l1.y_axis.delete = False
b1 += l1
db.add_chart(b1, 'B13')

# Pareto by failure type
b2 = BarChart(); b2.type = 'col'
b2.add_data(Reference(wk, min_col=10, min_row=25, max_row=35), titles_from_data=True)
b2.set_categories(Reference(wk, min_col=8, min_row=26, max_row=35))
style_bar(b2, 'Downtime by failure type (hrs)', 'Hours')
b2.series[0].graphicalProperties.solidFill = BLACK
l2 = LineChart()
l2.add_data(Reference(wk, min_col=11, min_row=25, max_row=35), titles_from_data=True)
l2.y_axis.axId = 200; l2.y_axis.number_format = '0%'; l2.y_axis.scaling.min = 0; l2.y_axis.scaling.max = 1
l2.y_axis.majorGridlines = None; l2.y_axis.crosses = 'max'; l2.y_axis.title = 'Running total'
l2.series[0].graphicalProperties.line.solidFill = GOLD
l2.series[0].graphicalProperties.line.width = 22000
l2.series[0].smooth = False
b2.x_axis.delete = False; b2.y_axis.delete = False; l2.y_axis.delete = False
b2 += l2
db.add_chart(b2, 'I13')

# Monthly trend
lt = BarChart(); lt.type = 'col'
lt.add_data(Reference(wk, min_col=4, min_row=39, max_row=51), titles_from_data=True)
lt.set_categories(Reference(wk, min_col=5, min_row=40, max_row=51))
style_bar(lt, 'Hours down by month (is it getting better?)', 'Hours')
lt.x_axis.delete = False; lt.y_axis.delete = False
lt.series[0].graphicalProperties.solidFill = GOLD
lt.series[0].graphicalProperties.line.solidFill = GOLD
db.add_chart(lt, 'B29')

# By shift
bs = BarChart(); bs.type = 'col'
bs.add_data(Reference(wk, min_col=9, min_row=39, max_row=44), titles_from_data=True)
bs.set_categories(Reference(wk, min_col=8, min_row=40, max_row=44))
style_bar(bs, 'Breakdowns by shift', 'Breakdowns')
bs.series[0].graphicalProperties.solidFill = BLACK
bs.x_axis.delete = False; bs.y_axis.delete = False
db.add_chart(bs, 'I29')

# Worst-first table under charts
r0 = 45
sect(db, r0, 'B', 'O', 'WORST MACHINES  (attack the top one first with 5 Whys or the Kaizen XP RCFA Template)')
hdr = [('B', 'D', 'MACHINE'), ('E', 'F', 'BREAKDOWNS'), ('G', 'H', 'HOURS DOWN'), ('I', 'J', 'MTTR (MINS)'), ('K', 'L', 'SHARE'), ('M', 'O', 'COST')]
for a, b, h in hdr:
    db.merge_cells(f'{a}{r0 + 1}:{b}{r0 + 1}'); db[f'{a}{r0 + 1}'] = h
    db[f'{a}{r0 + 1}'].font = Font(name=F, size=8, bold=True, color=INK)
    db[f'{a}{r0 + 1}'].alignment = Alignment(horizontal='center')
for i in range(5):
    r = r0 + 2 + i
    w = 7 + i
    vals = [('B', 'D', f'=Workings!H{w}', 'General'), ('E', 'F', f'=Workings!L{w}', '#,##0'), ('G', 'H', f'=Workings!J{w}', '0.0'),
            ('I', 'J', f'=IF(Workings!L{w}=0,"-",Workings!I{w}/Workings!L{w})', '#,##0'),
            ('K', 'L', f'=IF(Workings!C57=0,0,Workings!I{w}/Workings!C57)', '0%'),
            ('M', 'O', f'=Workings!I{w}/60*N(Setup!$C$11)', GBP)]
    for a, b, f_, fmt in vals:
        db.merge_cells(f'{a}{r}:{b}{r}')
        db[f'{a}{r}'] = f_; db[f'{a}{r}'].number_format = fmt
        db[f'{a}{r}'].font = Font(name=F, size=10, bold=(i == 0), color=INK)
        db[f'{a}{r}'].alignment = Alignment(horizontal='left' if a == 'B' else 'center', indent=1 if a == 'B' else 0)
        for cc in range(ord(a), ord(b) + 1):
            db[f'{chr(cc)}{r}'].border = box
            if i == 0:
                db[f'{chr(cc)}{r}'].fill = fill(CREAM)
db[f'B{r0 + 8}'] = '© 2026 Kaizen XP. Licensed for one person\'s use. Please do not share or resell.'
db[f'B{r0 + 8}'].font = Font(name=F, size=8, color='999999')
db.protection.sheet = True
db.sheet_properties.tabColor = GOLD
db.page_setup.orientation = 'landscape'
db.page_setup.fitToWidth = 1; db.page_setup.fitToHeight = 1
db.sheet_properties.pageSetUpPr.fitToPage = True
for ws_ in (st,):
    ws_.page_setup.fitToWidth = 1; ws_.page_setup.fitToHeight = 1; ws_.sheet_properties.pageSetUpPr.fitToPage = True

wb.active = 1
wb.properties.title = 'Breakdown Analysis Dashboard'
wb.properties.creator = 'Kaizen XP'
wb.save(OUT)
print('saved')
